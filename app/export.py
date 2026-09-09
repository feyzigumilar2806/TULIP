# ============================================================
# API DOWNLOAD DATA EXCEL TULIP
# ============================================================

from datetime import date
from datetime import datetime
from datetime import time
from datetime import timezone
from io import BytesIO
from typing import Annotated
from urllib.parse import urlencode
from zoneinfo import ZoneInfo

from fastapi import APIRouter
from fastapi import Depends
from fastapi import HTTPException
from fastapi import Query
from fastapi import Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.styles import Border
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.styles import Side
from sqlalchemy import select
from sqlalchemy.orm import Session
from sqlalchemy.orm import joinedload

from app.auth import get_current_user
from app.database import get_database
from app.models import Transaction
from app.models import User
from app.models import UserRole


# ============================================================
# PENGATURAN ROUTER
# ============================================================

router = APIRouter(
    prefix="/api/export",
    tags=["Download Excel"]
)

JAKARTA_TIMEZONE = ZoneInfo(
    "Asia/Jakarta"
)


# ============================================================
# MENGUBAH FILTER TANGGAL MENJADI UTC
# ============================================================

def start_of_date(
    selected_date: date
) -> datetime:

    jakarta_datetime = datetime.combine(
        selected_date,
        time.min,
        tzinfo=JAKARTA_TIMEZONE
    )

    return jakarta_datetime.astimezone(
        timezone.utc
    )


def end_of_date(
    selected_date: date
) -> datetime:

    jakarta_datetime = datetime.combine(
        selected_date,
        time.max,
        tzinfo=JAKARTA_TIMEZONE
    )

    return jakarta_datetime.astimezone(
        timezone.utc
    )


# ============================================================
# MENAMPILKAN WAKTU JAKARTA
# ============================================================

def format_jakarta_datetime(
    value: datetime | None
) -> str:

    if value is None:
        return ""

    if value.tzinfo is None:
        value = value.replace(
            tzinfo=timezone.utc
        )

    jakarta_time = value.astimezone(
        JAKARTA_TIMEZONE
    )

    return jakarta_time.strftime(
        "%d/%m/%Y %H:%M:%S"
    )


# ============================================================
# MENGAMANKAN TEKS EXCEL
# ============================================================

def safe_excel_text(
    value
) -> str:

    if value is None:
        return ""

    text_value = str(value)

    if text_value.startswith(
        ("=", "+", "-", "@")
    ):
        return "'" + text_value

    return text_value


# ============================================================
# MEMBUAT LINK BUKTI
# ============================================================

def create_proof_link(
    base_url: str,
    transaction_id: str,
    proof_type: str
) -> str:

    parameters = urlencode({
        "transaction_id":
            transaction_id,

        "proof_type":
            proof_type
    })

    return (
        f"{base_url}/?"
        f"{parameters}"
    )


# ============================================================
# MEMBERI TAMPILAN PADA HEADER
# ============================================================

def style_excel_header(
    worksheet
) -> None:

    yellow_fill = PatternFill(
        fill_type="solid",
        fgColor="FFD400"
    )

    header_font = Font(
        color="111111",
        bold=True
    )

    thin_side = Side(
        style="thin",
        color="B7B7B7"
    )

    header_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side
    )

    for cell in worksheet[1]:
        cell.fill = yellow_fill
        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        cell.border = header_border

    worksheet.row_dimensions[1].height = 35
    worksheet.freeze_panes = "A2"

    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )


# ============================================================
# MENGATUR LEBAR KOLOM
# ============================================================

def set_column_widths(
    worksheet
) -> None:

    column_widths = {
        "A": 29,
        "B": 22,
        "C": 18,
        "D": 28,
        "E": 20,
        "F": 24,
        "G": 22,
        "H": 14,
        "I": 14,
        "J": 14,
        "K": 14,
        "L": 20,
        "M": 28,
        "N": 24,
        "O": 30,
        "P": 24,
        "Q": 30,
        "R": 35,
    }

    for column, width in (
        column_widths.items()
    ):
        worksheet.column_dimensions[
            column
        ].width = width


# ============================================================
# MEMBERI TAMPILAN PADA LINK BUKTI
# ============================================================

def set_proof_hyperlink(
    cell,
    link: str
) -> None:

    cell.value = "Lihat Bukti"
    cell.hyperlink = link

    cell.font = Font(
        color="0563C1",
        underline="single",
        bold=True
    )

    cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )


# ============================================================
# MEMBUAT FILE EXCEL TRANSAKSI
# ============================================================

def create_excel_file(
    transactions: list[Transaction],
    base_url: str
) -> BytesIO:

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Data Transaksi"

    headers = [
        "ID Transaksi",
        "Tanggal dan Waktu",
        "Cabang",
        "Nama Karyawan",
        "NIK",
        "Unit Kerja",
        "Sub-unit Operation",
        "Jumlah Rp100",
        "Jumlah Rp200",
        "Jumlah Rp500",
        "Jumlah Rp1.000",
        "Total Nominal",
        "Status",
        "Bukti Penyerahan",
        "Nama Bukti Penyerahan",
        "Bukti Transfer",
        "Nama Bukti Transfer",
        "Catatan",
    ]

    worksheet.append(
        headers
    )

    for transaction in transactions:
        worksheet.append(
            [
                safe_excel_text(
                    transaction.id
                ),

                format_jakarta_datetime(
                    transaction.timestamp
                ),

                safe_excel_text(
                    transaction.branch.name
                    if transaction.branch
                    else ""
                ),

                safe_excel_text(
                    transaction.employee_name
                ),

                safe_excel_text(
                    transaction.employee_nik
                ),

                safe_excel_text(
                    transaction.unit
                ),

                safe_excel_text(
                    transaction.operation_subunit
                    or ""
                ),

                transaction.qty100,
                transaction.qty200,
                transaction.qty500,
                transaction.qty1000,
                transaction.total_amount,

                safe_excel_text(
                    transaction.status.value
                ),

                (
                    "Tersedia"
                    if transaction.handover_proof_path
                    else "Tidak tersedia"
                ),

                safe_excel_text(
                    transaction.handover_proof_name
                    or ""
                ),

                (
                    "Tersedia"
                    if transaction.transfer_proof_path
                    else "Tidak tersedia"
                ),

                safe_excel_text(
                    transaction.transfer_proof_name
                    or ""
                ),

                safe_excel_text(
                    transaction.notes
                    or ""
                ),
            ]
        )

        row_number = (
            worksheet.max_row
        )

        if transaction.handover_proof_path:
            handover_link = create_proof_link(
                base_url=base_url,
                transaction_id=transaction.id,
                proof_type="handover"
            )

            set_proof_hyperlink(
                worksheet.cell(
                    row=row_number,
                    column=14
                ),
                handover_link
            )

        if transaction.transfer_proof_path:
            transfer_link = create_proof_link(
                base_url=base_url,
                transaction_id=transaction.id,
                proof_type="transfer"
            )

            set_proof_hyperlink(
                worksheet.cell(
                    row=row_number,
                    column=16
                ),
                transfer_link
            )

    style_excel_header(
        worksheet
    )

    set_column_widths(
        worksheet
    )

    thin_side = Side(
        style="thin",
        color="D9D9D9"
    )

    data_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side
    )

    for row in worksheet.iter_rows(
        min_row=2
    ):
        for cell in row:
            cell.border = data_border

            if cell.hyperlink is None:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True
                )

    for row_number in range(
        2,
        worksheet.max_row + 1
    ):
        worksheet.cell(
            row=row_number,
            column=12
        ).number_format = (
            '"Rp" #,##0'
        )

    worksheet.sheet_view.showGridLines = (
        False
    )

    output = BytesIO()

    workbook.save(
        output
    )

    output.seek(0)

    return output


# ============================================================
# DOWNLOAD EXCEL
# ============================================================

@router.get("/transactions")
def download_transactions_excel(
    request: Request,

    current_user: Annotated[
        User,
        Depends(get_current_user)
    ],

    database: Annotated[
        Session,
        Depends(get_database)
    ],

    start_date: date | None = Query(
        default=None
    ),

    end_date: date | None = Query(
        default=None
    ),

    employee_name: str | None = Query(
        default=None,
        max_length=200
    ),

    branch_id: int | None = Query(
        default=None
    )
):
    if (
        start_date
        and end_date
        and start_date > end_date
    ):
        raise HTTPException(
            status_code=400,
            detail=(
                "Tanggal mulai tidak boleh "
                "melewati tanggal akhir."
            )
        )

    query = (
        select(Transaction)
        .options(
            joinedload(
                Transaction.branch
            )
        )
        .order_by(
            Transaction.timestamp.desc()
        )
    )

    # Kepala Cabang hanya dapat mengunduh
    # transaksi cabangnya sendiri.

    if (
        current_user.role
        == UserRole.BRANCH_HEAD
    ):
        if current_user.branch_id is None:
            raise HTTPException(
                status_code=403,
                detail=(
                    "Akun tidak memiliki "
                    "akses cabang."
                )
            )

        query = query.where(
            Transaction.branch_id
            == current_user.branch_id
        )

    # Pusat dan Superadmin dapat memilih cabang.

    elif branch_id is not None:
        query = query.where(
            Transaction.branch_id
            == branch_id
        )

    if start_date is not None:
        query = query.where(
            Transaction.timestamp
            >= start_of_date(
                start_date
            )
        )

    if end_date is not None:
        query = query.where(
            Transaction.timestamp
            <= end_of_date(
                end_date
            )
        )

    if employee_name:
        cleaned_employee_name = (
            employee_name.strip()
        )

        if cleaned_employee_name:
            query = query.where(
                Transaction.employee_name.ilike(
                    f"%{cleaned_employee_name}%"
                )
            )

    transactions = list(
        database.scalars(
            query
        ).unique().all()
    )

    base_url = str(
        request.base_url
    ).rstrip("/")

    excel_file = create_excel_file(
        transactions=transactions,
        base_url=base_url
    )

    current_time = datetime.now(
        JAKARTA_TIMEZONE
    ).strftime(
        "%Y%m%d-%H%M%S"
    )

    file_name = (
        f"Data-Transaksi-TULIP-"
        f"{current_time}.xlsx"
    )

    return StreamingResponse(
        excel_file,
        media_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="{file_name}"'
            )
        }
    )